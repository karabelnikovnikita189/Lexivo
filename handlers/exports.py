import os

from aiogram import Router, F
from docx import Document


from aiogram import types, Bot
from aiogram.types import FSInputFile

from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer
)

from reportlab.lib.styles import (
    getSampleStyleSheet
)

from database import get_last_transcription

router = Router()

from database import (
    get_last_transcription,
    is_user_vip
)


@router.callback_query(F.data == "export_pdf")
async def export_pdf(callback: types.CallbackQuery, bot: Bot):
    user_id = callback.from_user.id
    print("PDF USER:", user_id)
    print("IS VIP:", is_user_vip(user_id))


    if not is_user_vip(user_id):
        await callback.message.answer(
            "💎 PDF экспорт доступен только пользователям Lexivo PRO.\n\n"
            "Откройте раздел «✨ Подписка» для получения доступа."
        )
        return

    row = get_last_transcription(user_id)

    if not row:
        await callback.answer(
            "Нет данных для экспорта",
            show_alert=True
        )
        return

    data = {
        "title": row[1],
        "transcription": row[2],
        "summary": row[3],
        "translated": row[4]
    }

    export_text = (
        data.get("translated")
        or data.get("summary")
        or ""
    )

    filename = f"lexivo_{user_id}.pdf"

    doc = SimpleDocTemplate(filename)

    styles = getSampleStyleSheet()
    styles["Title"].fontName = "NotoSans"
    styles["Heading1"].fontName = "NotoSans"
    styles["BodyText"].fontName = "NotoSans"

    story = []

    story.append(
        Paragraph(
            data.get("title", "Lexivo Export"),
            styles["Title"]
        )
    )

    story.append(
        Spacer(1, 12)
    )

    if not data.get("translated"):
        story.append(
            Paragraph("Расшифровка", styles["Heading1"])
        )

        story.append(
            Paragraph(
                data["transcription"].replace("\n", "<br/>"),
                styles["BodyText"]
            )
        )


    story.append(
        Spacer(1, 12)
    )

    story.append(
        Paragraph("Анализ", styles["Heading1"])
    )

    story.append(
        Paragraph(
            export_text.replace("\n", "<br/>"),
            styles["BodyText"]
        )
    )

    doc.build(story)

    await bot.send_document(
        chat_id=user_id,
        document=FSInputFile(filename),
        caption="📄 PDF-документ готов"
    )

    os.remove(filename)

    await callback.answer()


@router.callback_query(F.data == "export_docx")
async def export_docx(callback: types.CallbackQuery,
                      bot: Bot):
    user_id = callback.from_user.id

    if not is_user_vip(user_id):
        await callback.message.answer(
            "💎 Word экспорт доступен только пользователям Lexivo PRO.\n\n"
            "Откройте раздел «✨ Подписка» для получения доступа."
        )
        return

    row = get_last_transcription(user_id)

    if not row:
        await callback.answer(
            "Нет данных для экспорта",
            show_alert=True
        )
        return

    data = {
        "title": row[1],
        "transcription": row[2],
        "summary": row[3],
        "translated": row[4]
    }

    filename = f"lexivo_{user_id}.docx"

    # Достаем перевод, а если его нет — берем оригинальный summary
    export_text = (
        data.get("translated")
        or data.get("summary")
        or ""
    )

    doc = Document()

    doc.add_heading(
        data.get("title", "Lexivo Export"),
        level=1
    )


    # Если перевода не было, экспортируем оригинальную транскрипцию
    if not data.get("translated"):

        doc.add_heading(
            "Расшифровка",
            level=2
        )
        doc.add_paragraph(
            data.get("transcription", "")
        )

    doc.add_heading(
        "Результат / Анализ",
        level=2
    )

    doc.add_paragraph(export_text)

    doc.save(filename)

    await bot.send_document(
        chat_id=user_id,
        document=FSInputFile(filename),
        caption="📝 Word-документ готов"
    )

    if os.path.exists(filename):
        os.remove(filename)

    await callback.answer()